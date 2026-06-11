import openpyxl
from openpyxl.styles import PatternFill, Font, Alignment

COLORS = {
    'Família': 'C6EFCE',
    'Original': 'BDD7EE',
    'Adultos': 'FFB3C6',
}

# Read DB cards
db_cards = []
with open(r'C:\Users\Barros Moreira\Desktop\Jogos\Cartas Radioativas\cartas-radioativas\BD_BKP\maldicao_db.txt', encoding='utf-8') as f:
    for line in f.readlines()[1:]:
        parts = line.strip().split('\t')
        if len(parts) >= 2 and parts[0]:
            db_cards.append((parts[0], parts[1], 'Banco de Dados'))

# Additional cards from PDFs and web
extra_cards = [
    # === Funny PDF (Família) ===
    ('Ter visão no escuro', 'Família', 'PDF Funny Scenarios'),
    ('Nunca mais ficar cansado', 'Família', 'PDF Funny Scenarios'),
    ('Ter cabeça de abóbora', 'Família', 'PDF Funny Scenarios'),
    ('Comer uma torta enorme de abóbora', 'Família', 'PDF Funny Scenarios'),
    ('Usar os sapatos no pé errado para sempre', 'Família', 'PDF Funny Scenarios'),
    ('Ter a língua presa em um poste congelado', 'Família', 'PDF Funny Scenarios'),
    ('Falar com sapos', 'Família', 'PDF Funny Scenarios'),
    ('Falar com cobras', 'Família', 'PDF Funny Scenarios'),
    ('Ter orelhas de duende', 'Família', 'PDF Funny Scenarios'),
    ('Ter barba de Papai Noel', 'Família', 'PDF Funny Scenarios'),
    ('Ter pelo em todo o corpo', 'Família', 'PDF Funny Scenarios'),
    ('Ter um amigo dinossauro', 'Família', 'PDF Funny Scenarios'),
    ('Ser um dinossauro', 'Família', 'PDF Funny Scenarios'),
    ('Cantar como Lady Gaga', 'Família', 'PDF Funny Scenarios'),
    ('Dançar como Mick Jagger', 'Família', 'PDF Funny Scenarios'),
    ('Andar no lombo de um elefante', 'Família', 'PDF Funny Scenarios'),
    ('Andar no lombo de um hipopótamo', 'Família', 'PDF Funny Scenarios'),
    ('Ter mãos enormes', 'Família', 'PDF Funny Scenarios'),
    ('Ter pés muito pequenos', 'Família', 'PDF Funny Scenarios'),
    ('Tomar banho em milkshake', 'Família', 'PDF Funny Scenarios'),
    ('Tomar banho em pudim de chocolate', 'Família', 'PDF Funny Scenarios'),
    ('Lamber um sapo', 'Família', 'PDF Funny Scenarios'),
    ('Comer uma larva', 'Família', 'PDF Funny Scenarios'),
    ('Visitar Hogwarts', 'Família', 'PDF Funny Scenarios'),
    ('Conhecer Jon Snow', 'Família', 'PDF Funny Scenarios'),
    ('Ter soluços por um longo tempo', 'Família', 'PDF Funny Scenarios'),
    ('Ter coceira em um lugar que não consegue coçar para sempre', 'Família', 'PDF Funny Scenarios'),
    ('Comer ração de cachorro', 'Família', 'PDF Funny Scenarios'),
    ('Comer uma maçã podre', 'Família', 'PDF Funny Scenarios'),
    ('Usar roupas de inverno no verão', 'Família', 'PDF Funny Scenarios'),
    ('Usar roupas de verão no inverno', 'Família', 'PDF Funny Scenarios'),
    ('Ter chifre de unicórnio', 'Família', 'PDF Funny Scenarios'),
    ('Ter nariz brilhante de rena', 'Família', 'PDF Funny Scenarios'),
    ('Beber um copo de sabão líquido', 'Família', 'PDF Funny Scenarios'),
    ('Beber um copo de vinagre', 'Família', 'PDF Funny Scenarios'),
    ('Soltar um pum alto na frente dos amigos', 'Família', 'PDF Funny Scenarios'),
    ('Fazer xixi na calça na frente dos amigos', 'Família', 'PDF Funny Scenarios'),
    ('Ter cabelo verde', 'Família', 'PDF Funny Scenarios'),
    ('Ter olhos roxos', 'Família', 'PDF Funny Scenarios'),
    ('Conhecer um alienígena', 'Família', 'PDF Funny Scenarios'),
    ('Comer pão com manteiga preta', 'Família', 'PDF Funny Scenarios'),
    ('Beber chá com areia', 'Família', 'PDF Funny Scenarios'),
    ('Morar em uma nave espacial', 'Família', 'PDF Funny Scenarios'),
    ('Nadar em um rio cheio de cobras', 'Família', 'PDF Funny Scenarios'),
    ('Deitar em uma cama cheia de aranhas', 'Família', 'PDF Funny Scenarios'),
    ('Ter um dragão de estimação', 'Família', 'PDF Funny Scenarios'),
    ('Beijar uma água-viva', 'Família', 'PDF Funny Scenarios'),
    ('Pisar em um ouriço', 'Família', 'PDF Funny Scenarios'),
    ('Ser o mais bonito da turma', 'Família', 'PDF Funny Scenarios'),
    ('Conseguir respirar debaixo d\'água', 'Família', 'PDF Funny Scenarios'),
    ('Criar um novo feriado', 'Família', 'PDF Funny Scenarios'),
    ('Criar um novo esporte', 'Família', 'PDF Funny Scenarios'),
    ('Cheirar o próprio pum', 'Família', 'PDF Funny Scenarios'),
    ('Cheirar o pum de outra pessoa', 'Família', 'PDF Funny Scenarios'),
    ('Ser um golfinho', 'Família', 'PDF Funny Scenarios'),
    ('Ser uma chita', 'Família', 'PDF Funny Scenarios'),
    ('Passar as festas com super-heróis', 'Família', 'PDF Funny Scenarios'),
    ('Ser uma estrela de rock quando crescer', 'Família', 'PDF Funny Scenarios'),
    ('Ganhar uma medalha olímpica de ouro', 'Família', 'PDF Funny Scenarios'),
    # === Kids1 PDF (Família) ===
    ('Ficar em quarentena com o Spongebob', 'Família', 'PDF Kids WYR'),
    ('Ficar em quarentena com o Harry Potter', 'Família', 'PDF Kids WYR'),
    ('Ter nariz de porco', 'Família', 'PDF Kids WYR'),
    ('Ter rosto de macaco', 'Família', 'PDF Kids WYR'),
    ('Ter um olho no meio da cabeça', 'Família', 'PDF Kids WYR'),
    ('Ter dois narizes', 'Família', 'PDF Kids WYR'),
    ('Ter mãos muito pequenas', 'Família', 'PDF Kids WYR'),
    ('Ter pés muito grandes', 'Família', 'PDF Kids WYR'),
    ('Usar shampoo para escovar os dentes', 'Família', 'PDF Kids WYR'),
    ('Usar pasta de dentes para lavar o cabelo', 'Família', 'PDF Kids WYR'),
    ('Passar a vida toda falando em rimas', 'Família', 'PDF Kids WYR'),
    ('Não falar nada por um ano inteiro', 'Família', 'PDF Kids WYR'),
    ('Comer comida que tem gosto bom mas parece cocô', 'Família', 'PDF Kids WYR'),
    ('Comer comida que parece boa mas tem gosto de cocô', 'Família', 'PDF Kids WYR'),
    ('Esfregar o chão do banheiro com a escova de dente por um dia inteiro', 'Família', 'PDF Kids WYR'),
    ('Usar a escova de dente que outra pessoa usou na própria boca', 'Família', 'PDF Kids WYR'),
    ('Usar roupas de desconhecidos diferentes todos os dias por uma semana', 'Família', 'PDF Kids WYR'),
    ('Soltar um pum muito alto na frente de todo mundo', 'Família', 'PDF Kids WYR'),
    ('Fazer cocô na calça silenciosamente', 'Família', 'PDF Kids WYR'),
    ('Crescer um par extra de orelhas mas ouvir mal', 'Família', 'PDF Kids WYR'),
    ('Crescer uma língua extra mas não sentir gosto', 'Família', 'PDF Kids WYR'),
    ('Usar gotas para os olhos feitas de vinagre', 'Família', 'PDF Kids WYR'),
    ('Usar papel higiênico feito de lixa', 'Família', 'PDF Kids WYR'),
    ('Tomar banho em cubos de gelo', 'Família', 'PDF Kids WYR'),
    ('Tomar banho em sopa de tomate', 'Família', 'PDF Kids WYR'),
    ('Ficar trancado em uma sala com 30 bebês chorando', 'Família', 'PDF Kids WYR'),
    ('Ter uma orelha de porco', 'Família', 'PDF Kids WYR'),
    ('Ter orelhas de elefante', 'Família', 'PDF Kids WYR'),
    ('Ter diarreia por duas semanas', 'Família', 'PDF Kids WYR'),
    ('Ficar em isolamento para sempre', 'Família', 'PDF Kids WYR'),
    ('Parecer magrinho mas ser forte', 'Família', 'PDF Kids WYR'),
    ('Parecer grandão mas ser fraco', 'Família', 'PDF Kids WYR'),
    # === Pick Your Poison PDF (Original) ===
    ('Fardar toda vez que beijar alguém', 'Original', 'PDF Pick Your Poison'),
    ('Lutar com uma avestruz até a morte', 'Original', 'PDF Pick Your Poison'),
    ('Sempre soar como se tivesse inalado hélio', 'Original', 'PDF Pick Your Poison'),
    ('Perder um dente por ano', 'Original', 'PDF Pick Your Poison'),
    ('Só poder se locomover andando a pé', 'Original', 'PDF Pick Your Poison'),
    ('Ser mortalmente alérgico à luz solar', 'Original', 'PDF Pick Your Poison'),
    ('Ter que cantar todas as suas falas ao invés de falar', 'Original', 'PDF Pick Your Poison'),
    ('Ter que dançar em vez de andar para qualquer lugar', 'Original', 'PDF Pick Your Poison'),
    ('Nunca poder dizer não', 'Original', 'PDF Pick Your Poison'),
    ('Nunca poder dizer sim', 'Original', 'PDF Pick Your Poison'),
    ('Ter flashbacks aleatórios de memórias embaraçosas toda hora', 'Original', 'PDF Pick Your Poison'),
    ('Tossir quando mentir', 'Original', 'PDF Pick Your Poison'),
    ('Espirrar quando sentir raiva', 'Original', 'PDF Pick Your Poison'),
    ('Ter que sussurrar tudo', 'Original', 'PDF Pick Your Poison'),
    ('Ter que gritar tudo', 'Original', 'PDF Pick Your Poison'),
    ('Andar de ré para todo lugar', 'Original', 'PDF Pick Your Poison'),
    ('Nunca poder sentar em uma cadeira', 'Original', 'PDF Pick Your Poison'),
    ('Ter que usar fantasia de palhaço no trabalho por um mês', 'Original', 'PDF Pick Your Poison'),
    ('Ter que usar pijama em todo lugar por um ano', 'Original', 'PDF Pick Your Poison'),
    ('Não conseguir distinguir salgado de doce', 'Original', 'PDF Pick Your Poison'),
    ('Sentir gosto de papelão em tudo que comer', 'Original', 'PDF Pick Your Poison'),
    ('Ter que fazer check-in nas redes sociais em todo lugar que for', 'Original', 'PDF Pick Your Poison'),
    ('Nunca mais poder usar óculos de sol', 'Original', 'PDF Pick Your Poison'),
    ('Ter que fazer uma reverência para todo mundo que encontrar', 'Original', 'PDF Pick Your Poison'),
    ('Só poder se comunicar por mímica', 'Original', 'PDF Pick Your Poison'),
    ('Ser obrigado a comprar qualquer coisa que alguém te oferecer', 'Original', 'PDF Pick Your Poison'),
    ('Esquecer o nome de todos que conhece após 24 horas', 'Original', 'PDF Pick Your Poison'),
    ('Nunca esquecer nenhuma briga que já teve', 'Original', 'PDF Pick Your Poison'),
    ('Ter que pedir permissão para ir ao banheiro a vida toda', 'Original', 'PDF Pick Your Poison'),
    # === Pick Your Poison NSFW (Adultos) ===
    ('Participar de uma orgia em uma casa de repouso', 'Adultos', 'PDF Pick Your Poison NSFW'),
    ('Matar e comer seu melhor amigo', 'Adultos', 'PDF Pick Your Poison NSFW'),
    ('Colocar a mão em um liquidificador e ligar', 'Adultos', 'PDF Pick Your Poison NSFW'),
    ('Fazer sexo com alguém que cheira muito mal', 'Adultos', 'PDF Pick Your Poison NSFW'),
    ('Ser filmado fazendo sexo sem saber e o vídeo vazar', 'Adultos', 'PDF Pick Your Poison NSFW'),
    ('Ter relações íntimas com alguém 40 anos mais velho', 'Adultos', 'PDF Pick Your Poison NSFW'),
    ('Nunca mais ter orgasmo', 'Adultos', 'PDF Pick Your Poison NSFW'),
    ('Ter orgasmo toda vez que espirra em público', 'Adultos', 'PDF Pick Your Poison NSFW'),
    ('Fazer sexo ao vivo em cadeia nacional de televisão', 'Adultos', 'PDF Pick Your Poison NSFW'),
    ('Ter que seduzir alguém que não te atrai para salvar sua vida', 'Adultos', 'PDF Pick Your Poison NSFW'),
    # === Web brightful.me (Original / Família) ===
    ('Viver em país com custo de vida baixo mas clima horrível', 'Original', 'brightful.me'),
    ('Ser muito inteligente mas totalmente antipático', 'Original', 'brightful.me'),
    ('Ter muitos amigos mas nenhum dinheiro', 'Original', 'brightful.me'),
    ('Ter muito dinheiro mas nenhum amigo verdadeiro', 'Original', 'brightful.me'),
    ('Saber a data de todas as mortes de quem você ama', 'Original', 'brightful.me'),
    ('Poder salvar uma vida mas perder algo importante para você', 'Original', 'brightful.me'),
    ('Descobrir que toda a sua vida foi uma mentira', 'Original', 'brightful.me'),
    ('Acordar todo dia sem saber quem você é por 10 minutos', 'Original', 'brightful.me'),
    ('Nunca poder ouvir sua música favorita de novo', 'Original', 'brightful.me'),
    ('Só poder ouvir sua música favorita e mais nada', 'Original', 'brightful.me'),
    ('Viver em uma simulação e descobrir isso', 'Original', 'brightful.me'),
    ('Saber a verdade sobre o universo mas não poder contar a ninguém', 'Original', 'brightful.me'),
    ('Perder sua personalidade mas ganhar muito sucesso', 'Original', 'brightful.me'),
    ('Manter sua personalidade mas nunca ter sucesso em nada', 'Original', 'brightful.me'),
    ('Ter poder ilimitado mas ninguém saber que foi você', 'Original', 'brightful.me'),
    ('Ser mundialmente famoso por algo que não fez', 'Original', 'brightful.me'),
    ('Ser esquecido por todos após 24 horas de cada encontro', 'Original', 'brightful.me'),
    ('Lembrar de tudo que alguém já disse para você', 'Original', 'brightful.me'),
    ('Ter que escolher entre salvar um estranho ou seu animal de estimação', 'Original', 'brightful.me'),
    ('Nunca mais poder mentir sobre nada', 'Original', 'brightful.me'),
    ('Saber quando alguém está mentindo mas não poder provar', 'Original', 'brightful.me'),
    ('Ser capaz de curar qualquer doença mas contrair todas elas primeiro', 'Original', 'brightful.me'),
    ('Morar sozinho em uma ilha com toda tecnologia do mundo', 'Original', 'brightful.me'),
    ('Morar em uma cidade movimentada sem poder usar nenhuma tecnologia', 'Original', 'brightful.me'),
    ('Viver 500 anos mas sozinho por metade do tempo', 'Original', 'brightful.me'),
    ('Ter tudo que sempre quis mas perder a capacidade de sentir alegria', 'Original', 'brightful.me'),
    ('Sentir alegria por qualquer coisa mas não conseguir nada que quer', 'Original', 'brightful.me'),
    ('Ter acesso a todo conhecimento do mundo mas esquecer tudo à meia-noite', 'Original', 'brightful.me'),
    ('Poder voar mas somente à velocidade de uma tartaruga', 'Família', 'brightful.me'),
    ('Poder teletransportar mas somente dentro do seu quarteirão', 'Família', 'brightful.me'),
    ('Ter super força mas somente para abrir potes', 'Família', 'brightful.me'),
    ('Ser invisível mas somente quando ninguém está olhando para você', 'Família', 'brightful.me'),
    # === gohen.com - Funny (Família) ===
    ('Ter que imitar alguém de trabalho a vida toda', 'Família', 'gohen.com'),
    ('Ter que usar um traje de gorila onde for', 'Família', 'gohen.com'),
    ('Sempre acordar com o cabelo num penteado ridículo', 'Família', 'gohen.com'),
    ('Só poder comer comida em formato de estrela', 'Família', 'gohen.com'),
    ('Ter que fazer uma dança a cada vez que entra em um aposento', 'Família', 'gohen.com'),
    ('Ter que aplaudir de pé cada vez que alguém entra no cômodo', 'Família', 'gohen.com'),
    ('Rir de forma hilária em situações sérias', 'Família', 'gohen.com'),
    ('Falar como um pirata por um ano inteiro', 'Família', 'gohen.com'),
    ('Roncar tão alto que acorda vizinhos', 'Família', 'gohen.com'),
    ('Ser seguido por uma câmera de reality show a vida toda', 'Família', 'gohen.com'),
    ('Ter que cumprimentar todos com um aperto de mão bizarro que você inventou', 'Família', 'gohen.com'),
    # === gohen.com - Hard (Original) ===
    ('Descobrir que foi adotado aos 30 anos', 'Original', 'gohen.com'),
    ('Descobrir que seu melhor amigo falou suas maiores mentiras para todos', 'Original', 'gohen.com'),
    ('Nunca descobrir quem foi que arruinou sua vida', 'Original', 'gohen.com'),
    ('Saber quem arruinou sua vida mas nunca poder fazer nada', 'Original', 'gohen.com'),
    ('Perder tudo que tem mas saber que foi por sua própria culpa', 'Original', 'gohen.com'),
    ('Perder tudo que tem e não saber por quê', 'Original', 'gohen.com'),
    ('Ter que abrir mão de 10 anos de vida para dar a alguém que ama', 'Original', 'gohen.com'),
    ('Ver a morte de um desconhecido que podia ter evitado', 'Original', 'gohen.com'),
    ('Ser responsável pela derrota do seu time nos finais do campeonato', 'Original', 'gohen.com'),
    ('Ser o único sobrevivente de um desastre', 'Original', 'gohen.com'),
    ('Trabalhar no emprego dos sonhos por 1 ano e depois nunca mais poder trabalhar nele', 'Original', 'gohen.com'),
    ('Trabalhar em um emprego horrível por 10 anos para poder se aposentar cedo', 'Original', 'gohen.com'),
    # === gohen.com - Dirty (Adultos) ===
    ('Fazer sexo com familiar assistindo sem saber', 'Adultos', 'gohen.com'),
    ('Pegar seus pais fazendo sexo', 'Adultos', 'gohen.com'),
    ('Ter relações com alguém famoso mas nunca poder contar a ninguém', 'Adultos', 'gohen.com'),
    ('Transar em um avião lotado', 'Adultos', 'gohen.com'),
    ('Nunca mais ter relações íntimas ou ter relações somente com desconhecidos', 'Adultos', 'gohen.com'),
    ('Ter que contar todos os detalhes da sua vida sexual para seus pais', 'Adultos', 'gohen.com'),
    ('Seus pais contarem todos os detalhes da vida sexual deles para você', 'Adultos', 'gohen.com'),
    ('Ser filmado sem saber em uma situação íntima', 'Adultos', 'gohen.com'),
    ('Descobrir que seu parceiro de uma noite é famoso depois que a noite acabou', 'Adultos', 'gohen.com'),
    ('Mandar uma mensagem íntima para o grupo da família por engano', 'Adultos', 'gohen.com'),
    # === bungie.net forum (Adultos) ===
    ('Comer seu próprio corpo até sobrar só a cabeça', 'Adultos', 'bungie.net'),
    ('Ter que matar seu animal de estimação com as próprias mãos para sobreviver', 'Adultos', 'bungie.net'),
    ('Ter que escolher qual dos seus filhos vai morrer', 'Adultos', 'bungie.net'),
    ('Ser torturado por um dia ou perder um membro', 'Adultos', 'bungie.net'),
    ('Acordar nu em uma cidade desconhecida sem documentos ou dinheiro', 'Adultos', 'bungie.net'),
    ('Saber que vai morrer em uma semana e não poder contar a ninguém', 'Adultos', 'bungie.net'),
    ('Ter que assistir um ente querido morrer sem poder fazer nada', 'Adultos', 'bungie.net'),
    ('Ser traído pelo seu melhor amigo da pior forma possível', 'Adultos', 'bungie.net'),
    ('Descobrir que causou a morte de alguém sem querer há anos', 'Adultos', 'bungie.net'),
    ('Perder os sentidos da visão e audição ao mesmo tempo', 'Adultos', 'bungie.net'),
    ('Ser preso por um crime que cometeu mas acha que foi justo', 'Adultos', 'bungie.net'),
    ('Ser preso por um crime que não cometeu e não poder provar sua inocência', 'Adultos', 'bungie.net'),
    ('Descobrir que toda a sua família sempre te odiou secretamente', 'Adultos', 'bungie.net'),
    ('Saber que seu filho vai ser uma pessoa terrível quando crescer', 'Adultos', 'bungie.net'),
    ('Viver em um mundo pós-apocalíptico sozinho para sempre', 'Adultos', 'bungie.net'),
    ('Perder toda memória dos melhores momentos da sua vida', 'Adultos', 'bungie.net'),
    ('Ter que mentir para salvar alguém que você ama mas destruir sua reputação', 'Adultos', 'bungie.net'),
    ('Nunca mais poder sonhar', 'Adultos', 'bungie.net'),
    ('Nunca mais poder chorar mesmo sentindo dor', 'Adultos', 'bungie.net'),
    ('Ser consciente mas paralítico por um mês', 'Adultos', 'bungie.net'),
]

all_cards = db_cards + extra_cards

# Sort: Família first, then Original, then Adultos; alphabetical within each
order = ['Família', 'Original', 'Adultos']
all_cards.sort(key=lambda x: (order.index(x[1]) if x[1] in order else 99, x[0].lower()))

# Create workbook
wb = openpyxl.Workbook()
ws = wb.active
ws.title = "Maldições"

header_fill = PatternFill(start_color='2D0040', end_color='2D0040', fill_type='solid')
header_font = Font(bold=True, color='FFFFFF', size=12)
center = Alignment(horizontal='center', vertical='center')

# Header row
for col, title in enumerate(['Texto', 'Categoria', 'Fonte'], 1):
    cell = ws.cell(row=1, column=col, value=title)
    cell.fill = header_fill
    cell.font = header_font
    cell.alignment = center

ws.row_dimensions[1].height = 22

# Data rows
for row_num, (texto, cat, fonte) in enumerate(all_cards, 2):
    ws.cell(row=row_num, column=1, value=texto)
    ws.cell(row=row_num, column=2, value=cat)
    ws.cell(row=row_num, column=3, value=fonte)
    color = COLORS.get(cat, 'FFFFFF')
    row_fill = PatternFill(start_color=color, end_color=color, fill_type='solid')
    for col in range(1, 4):
        cell = ws.cell(row=row_num, column=col)
        cell.fill = row_fill
        cell.alignment = Alignment(vertical='center', wrap_text=(col == 1))

# Column widths
ws.column_dimensions['A'].width = 75
ws.column_dimensions['B'].width = 12
ws.column_dimensions['C'].width = 28

# Freeze header
ws.freeze_panes = 'A2'

# Auto-filter
ws.auto_filter.ref = f'A1:C{len(all_cards)+1}'

# Summary sheet
ws2 = wb.create_sheet('Resumo')
ws2.column_dimensions['A'].width = 20
ws2.column_dimensions['B'].width = 12

summary_header = [('Categoria', 'Total')]
counts = {}
for _, cat, _ in all_cards:
    counts[cat] = counts.get(cat, 0) + 1

ws2.cell(row=1, column=1, value='Categoria').font = Font(bold=True)
ws2.cell(row=1, column=2, value='Total').font = Font(bold=True)
ws2.cell(row=1, column=3, value='Fonte').font = Font(bold=True)
ws2.column_dimensions['C'].width = 28

source_counts = {}
for _, cat, fonte in all_cards:
    key = (cat, fonte)
    source_counts[key] = source_counts.get(key, 0) + 1

r = 2
for cat in order:
    if cat in counts:
        cell = ws2.cell(row=r, column=1, value=cat)
        ws2.cell(row=r, column=2, value=counts[cat])
        color = COLORS.get(cat, 'FFFFFF')
        for c in range(1, 3):
            ws2.cell(row=r, column=c).fill = PatternFill(start_color=color, end_color=color, fill_type='solid')
        r += 1

r += 1
ws2.cell(row=r, column=1, value='TOTAL').font = Font(bold=True)
ws2.cell(row=r, column=2, value=len(all_cards)).font = Font(bold=True)
r += 2

ws2.cell(row=r, column=1, value='Por fonte:').font = Font(bold=True)
r += 1
src_totals = {}
for _, _, fonte in all_cards:
    src_totals[fonte] = src_totals.get(fonte, 0) + 1
for src, cnt in sorted(src_totals.items(), key=lambda x: -x[1]):
    ws2.cell(row=r, column=1, value=src)
    ws2.cell(row=r, column=2, value=cnt)
    r += 1

output = r'C:\Users\Barros Moreira\Desktop\Jogos\Maldição\maldicoes.xlsx'
wb.save(output)
print(f"Arquivo salvo: {output}")
print(f"Total de cartas: {len(all_cards)}")
for cat in order:
    print(f"  {cat}: {counts.get(cat, 0)}")
